from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Slide Layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]

    # --- SLIDE 1: Title Slide ---
    slide1 = prs.slides.add_slide(title_slide_layout)
    title1 = slide1.shapes.title
    subtitle1 = slide1.placeholders[1]
    
    title1.text = "ADR-Detect: Clinical Decision Support"
    subtitle1.text = (
        "Automated Adverse Drug Reaction Detection Using TF-IDF & Random Forest\n\n"
        "Student: Pranav Prabhakar Dabholkar | Roll No: [INSERT ROLL NUMBER]\n"
        "M.Sc. Data Science | KES' Shroff College of Arts & Commerce\n"
        "Subject: Advanced Text Analytics and Natural Language Processing\n"
        "Guide: [TEACHER NAME]"
    )

    # --- SLIDE 2: Executive Summary ---
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    slide2.shapes.title.text = "Executive Summary & Clinical Problem"
    tf2 = slide2.shapes.placeholders[1].text_frame
    tf2.text = "The Problem: Manual Review Bottleneck"
    
    p = tf2.add_paragraph()
    p.text = "Hospitals and pharmacovigilance bodies receive millions of unstructured clinical notes."
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "Manual triage of post-market reports is slow and prone to fatigue."
    p.level = 1
    
    p = tf2.add_paragraph()
    p.text = "The NLP Solution: ADR-Detect"
    p.level = 0
    p = tf2.add_paragraph()
    p.text = "Automated screening converts medical notes into bigram TF-IDF vectors."
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "Evaluated on the benchmark ADE-Corpus-V2 dataset (23,516 sentences)."
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "[!] Insert Class Distribution & Sentence Length EDA charts here."
    p.level = 1

    # --- SLIDE 3: Technical Blueprint ---
    slide3 = prs.slides.add_slide(bullet_slide_layout)
    slide3.shapes.title.text = "Technical Blueprint & Architecture"
    tf3 = slide3.shapes.placeholders[1].text_frame
    tf3.text = "Dataset & Environment"
    
    p = tf3.add_paragraph()
    p.text = "23,516 Labeled Sentences | TF-IDF (1, 2) N-Gram Range | Python 3.10"
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "Pipeline Flow"
    p.level = 0
    p = tf3.add_paragraph()
    p.text = "1. Raw Clinical Text Ingestion"
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "2. Regex Cleaning (Standardization, Bracket Removal)"
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "3. TF-IDF Transformation"
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "4. Random Forest Classification"
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "5. Streamlit Deployment with joblib caching"
    p.level = 1

    # --- SLIDE 4: User Interface ---
    slide4 = prs.slides.add_slide(bullet_slide_layout)
    slide4.shapes.title.text = "User Interface (Dashboard)"
    tf4 = slide4.shapes.placeholders[1].text_frame
    tf4.text = "Streamlit Application Highlights"
    
    p = tf4.add_paragraph()
    p.text = "Zero-Friction Ingestion: Supports free-text clinical notes."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "One-Click Clinical Scenarios: Pre-loaded presets for rapid demonstration."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "Real-Time Model Information: Transparent baseline metrics displayed."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "[!] Insert Screenshot 1 (Empty/Initial Streamlit Dashboard) here."
    p.level = 0

    # --- SLIDE 5: Empirical Verification 1 ---
    slide5 = prs.slides.add_slide(bullet_slide_layout)
    slide5.shapes.title.text = "Empirical Verification: Adverse Reaction"
    tf5 = slide5.shapes.placeholders[1].text_frame
    tf5.text = "Case Study 1: Angioedema Detection"
    
    p = tf5.add_paragraph()
    p.text = "Clinical Input: \"The patient developed acute angioedema of the lips...\""
    p.level = 1
    p = tf5.add_paragraph()
    p.text = "Prediction: Adverse Reaction Detected (Class 1)"
    p.level = 1
    p = tf5.add_paragraph()
    p.text = "Signal Highlights: High-weight tokens like 'lips' and 'tongue' extracted."
    p.level = 1
    p = tf5.add_paragraph()
    p.text = "[!] Insert Screenshot 2 (Red Alert Card) here."
    p.level = 0

    # --- SLIDE 6: Empirical Verification 2 ---
    slide6 = prs.slides.add_slide(bullet_slide_layout)
    slide6.shapes.title.text = "Empirical Verification: Negative Control"
    tf6 = slide6.shapes.placeholders[1].text_frame
    tf6.text = "Case Study 2: Routine Follow-up"
    
    p = tf6.add_paragraph()
    p.text = "Clinical Input: \"The patient reports feeling well and tolerating...\""
    p.level = 1
    p = tf6.add_paragraph()
    p.text = "Prediction: No Adverse Reaction Detected (Class 0)"
    p.level = 1
    p = tf6.add_paragraph()
    p.text = "Signal Highlights: Neutral non-adverse tokens identified (e.g., 'complaints', 'regimen')."
    p.level = 1
    p = tf6.add_paragraph()
    p.text = "[!] Insert Screenshot 3 (Green Confirmation Card) here."
    p.level = 0

    # --- SLIDE 7: Evaluation Metrics ---
    slide7 = prs.slides.add_slide(bullet_slide_layout)
    slide7.shapes.title.text = "Model Evaluation & Feature Matrix"
    tf7 = slide7.shapes.placeholders[1].text_frame
    tf7.text = "Performance on 4,704 Unseen Sentences"
    
    p = tf7.add_paragraph()
    p.text = "Overall Accuracy: 89%"
    p.level = 1
    p = tf7.add_paragraph()
    p.text = "Class 1 (Adverse Reaction) Precision: 86%"
    p.level = 1
    p = tf7.add_paragraph()
    p.text = "Class 1 (Adverse Reaction) Recall: 74%"
    p.level = 1
    p = tf7.add_paragraph()
    p.text = "Class 0 (No Reaction) Recall: 95%"
    p.level = 1
    p = tf7.add_paragraph()
    p.text = "[!] Insert Screenshot 6 (Confusion Matrix) here."
    p.level = 0

    # --- SLIDE 8: Summary & Future Scope ---
    slide8 = prs.slides.add_slide(bullet_slide_layout)
    slide8.shapes.title.text = "Summary & Future Scope"
    tf8 = slide8.shapes.placeholders[1].text_frame
    tf8.text = "Key Achievements"
    
    p = tf8.add_paragraph()
    p.text = "Engineered an end-to-end NLP triage engine from raw clinical corpora."
    p.level = 1
    p = tf8.add_paragraph()
    p.text = "Built a self-contained, interactive Streamlit dashboard."
    p.level = 1
    
    p = tf8.add_paragraph()
    p.text = "Future Enhancements"
    p.level = 0
    p = tf8.add_paragraph()
    p.text = "Named Entity Recognition (NER): Transition to token-level entity extraction."
    p.level = 1
    p = tf8.add_paragraph()
    p.text = "Transformer Fine-Tuning: Benchmark against BioBERT to improve recall."
    p.level = 1

    # Save the presentation
    prs.save('ADR_Detect_Presentation.pptx')
    print("Presentation generated successfully as 'ADR_Detect_Presentation.pptx'")

if __name__ == '__main__':
    create_presentation()